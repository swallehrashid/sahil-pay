"""Backup & Delete: owner-only, password + exact phrase, full backup first, then nothing left."""

import io

import pytest
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from models import (
    AuditLog, ChargeCategory, Invoice, Payment, Property, TeamMember, Tenant, Unit,
)
from services.category_service import seed_default_categories
from tests.test_access_control import (
    _auth, _make_invoice, _make_landlord, _make_payment, _make_property, _make_team_member,
    _make_tenant, _make_unit, _token,
)


@pytest.fixture()
def client(app):
    return app.test_client()


@pytest.fixture()
def account(app, db_session):
    landlord, user = _make_landlord(db_session, "wipe")
    user.password_hash = generate_password_hash("RealEstate254")
    seed_default_categories(landlord.id)
    for b in range(2):
        prop = _make_property(db_session, landlord, f"W{b}")
        for u in range(3):
            unit = _make_unit(db_session, prop, f"U{u}")
            tenant = _make_tenant(db_session, landlord, unit)
            _make_invoice(db_session, landlord, tenant, prop, unit)
            _make_payment(db_session, landlord, tenant, prop, unit)
    tm, tm_user = _make_team_member(db_session, landlord, property_ids=[prop.id])
    other, other_user = _make_landlord(db_session, "keep")
    other_prop = _make_property(db_session, other, "Other")
    _make_tenant(db_session, other, _make_unit(db_session, other_prop, "O1"))
    db_session.flush()
    return {
        "landlord": landlord, "owner_h": _auth(_token(app, user, landlord_id=landlord.id)),
        "team_h": _auth(_token(app, tm_user, landlord_id=landlord.id, team_member_id=tm.id)),
        "tm": tm, "other": other,
    }


def _token_for(client, h, password="RealEstate254"):
    return client.post("/api/settings/wipe/verify-password", headers=h, json={"password": password})


def test_team_members_can_never_wipe(client, account):
    h = account["team_h"]
    assert client.get("/api/settings/wipe", headers=h).status_code == 403
    assert _token_for(client, h).status_code == 403
    assert client.post("/api/settings/wipe", headers=h, json={}).status_code == 403


def test_wrong_password_and_wrong_phrase_are_refused(client, account):
    h = account["owner_h"]
    assert _token_for(client, h, "nope").status_code == 403
    token = _token_for(client, h).get_json()["wipe_token"]
    r = client.post("/api/settings/wipe", headers=h,
                    json={"wipe_token": token, "phrase": "delete account", "confirm": True})
    assert r.status_code == 400
    r = client.post("/api/settings/wipe", headers=h, json={"wipe_token": "forged", "phrase": "x", "confirm": True})
    assert r.status_code == 403
    assert Tenant.query.filter_by(landlord_id=account["landlord"].id).count() == 6


def test_full_flow_backs_up_then_leaves_nothing(client, account):
    ll, h = account["landlord"], account["owner_h"]
    info = client.get("/api/settings/wipe", headers=h).get_json()
    assert info["counts"]["tenants"] == 6 and info["phrase"] == f"DELETE ACCOUNT SP-{ll.id:05d}"

    token = _token_for(client, h).get_json()["wipe_token"]
    r = client.post("/api/settings/wipe", headers=h,
                    json={"wipe_token": token, "phrase": info["phrase"], "confirm": True})
    assert r.status_code == 200, r.get_json()
    body = r.get_json()
    assert body["deleted"]["tenants"] == 6
    assert all(v == 0 for v in body["remaining"].values())

    for model in (Property, Tenant, Invoice, Payment):
        assert model.query.filter_by(landlord_id=ll.id).count() == 0
    assert Unit.query.join(Property).filter(Property.landlord_id == ll.id).count() == 0

    # What stays: the account, its team, its categories, the audit trail — and other accounts.
    assert TeamMember.query.filter_by(landlord_id=ll.id).count() == 1
    assert ChargeCategory.query.filter_by(landlord_id=ll.id).count() >= 1
    assert AuditLog.query.filter_by(landlord_id=ll.id, action="backup_and_delete_account").count() == 1
    assert Tenant.query.filter_by(landlord_id=account["other"].id).count() == 1

    # The backup holds what was deleted.
    dl = client.get(body["download_url"], headers=h)
    assert dl.status_code == 200
    wb = load_workbook(io.BytesIO(dl.data), read_only=True)
    assert {"Properties", "Units", "Tenants", "Invoices", "Payments"} <= set(wb.sheetnames)
    assert wb["Tenants"].max_row == 7   # header + 6
    assert wb["Units"].max_row == 7
    assert wb["Properties"].max_row == 3

    # A team member cannot fetch the backup.
    assert client.get(body["download_url"], headers=account["team_h"]).status_code == 403
